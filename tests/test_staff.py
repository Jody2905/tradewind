"""Staff pages, refunds, the sales and VAT report, and its Excel export."""

import io
from datetime import date, datetime, timedelta, timezone as tz

import pytest
from django.db import IntegrityError, transaction
from django.urls import reverse
from openpyxl import load_workbook

from shop import exports, reports, services
from shop.models import Order, Product

from .conftest import NOW, items

STAFF_PAGES = ["staff_dashboard", "staff_orders", "staff_stock", "staff_sales"]


@pytest.fixture
def staff_client(client, staff):
    client.force_login(staff)
    return client


@pytest.fixture
def order(customer, products):
    return services.place_order(customer, items(products, hot_sauce=2, basket=1), "bank_transfer", "collect",
                                "784-555-0101", now=NOW)


# ---- Who can get in -----------------------------------------------------------------------

@pytest.mark.parametrize("page", STAFF_PAGES)
def test_staff_pages_need_login(client, db, page):
    r = client.get(reverse(page))
    assert r.status_code == 302 and reverse("login") in r["Location"]


@pytest.mark.parametrize("page", STAFF_PAGES)
def test_customers_get_not_found(client, customer, page):
    client.force_login(customer)
    assert client.get(reverse(page)).status_code == 404


def test_customers_cannot_post_staff_actions(client, customer, order):
    client.force_login(customer)
    r = client.post(reverse("staff_order_action", args=[order.number, "paid"]), {"reference": "fake"})
    assert r.status_code == 404 and Order.objects.get().status == "awaiting_payment"


@pytest.mark.parametrize("page", STAFF_PAGES)
def test_staff_can_open_every_page(staff_client, order, page):
    assert staff_client.get(reverse(page)).status_code == 200


# ---- Processing an order --------------------------------------------------------------------

def act(client, order, action, **data):
    return client.post(reverse("staff_order_action", args=[order.number, action]), data, follow=True)


def test_payment_packing_and_completion(staff_client, staff, customer, products):
    # placed "now", because the staff actions below happen at the real current time
    order = services.place_order(customer, items(products, basket=1), "bank_transfer", "collect", "784")
    assert "packing queue" in act(staff_client, order, "paid", reference="BT-1001").content.decode()
    assert staff_client.get(reverse("staff_slip", args=[order.number])).status_code == 200
    act(staff_client, order, "packed")
    html = act(staff_client, order, "completed").content.decode()
    assert "is complete" in html and Order.objects.get().status == "completed"
    steps = list(Order.objects.get().events.all())
    assert [e.to_status for e in steps] == ["awaiting_payment", "paid", "packed", "completed"]
    assert all(e.by == staff for e in steps[1:])                  # the history records who did each step


def test_wrong_step_shows_the_rule(staff_client, order):
    html = act(staff_client, order, "packed").content.decode()
    assert "can&#x27;t be marked packed" in html


def test_packing_slip_only_for_orders_being_packed(staff_client, order):
    assert staff_client.get(reverse("staff_slip", args=[order.number])).status_code == 404


def test_unknown_action_is_404(staff_client, order):
    assert staff_client.post(reverse("staff_order_action", args=[order.number, "delete"])).status_code == 404


def test_refund_after_paid_order_is_cancelled(staff_client, staff, order):
    services.mark_paid(staff, order, "BT-1001", now=NOW)
    act(staff_client, order, "cancel", reason="customer moved")
    assert Order.objects.get().refund_due
    assert "1 cancelled order still needs a refund" in staff_client.get(reverse("staff_dashboard")).content.decode()
    html = act(staff_client, order, "refund", reference="").content.decode()
    assert "refund reference" in html
    act(staff_client, order, "refund", reference="CHQ-3341")
    o = Order.objects.get()
    assert not o.refund_due and o.refund_reference == "CHQ-3341" and o.refunded_at
    with pytest.raises(services.ShopError, match="No refund is due"):
        services.record_refund(staff, o, "again")


def test_refund_can_only_follow_a_payment(order):
    with pytest.raises(IntegrityError), transaction.atomic():
        Order.objects.filter(pk=order.pk).update(refunded_at=NOW)          # never paid


def test_expire_button(staff_client, customer, products):
    services.place_order(customer, items(products, nutmeg=1), "collection", "collect", "784", now=NOW - timedelta(days=30))
    html = staff_client.post(reverse("staff_expire"), follow=True).content.decode()
    assert "1 overdue unpaid order(s) cancelled" in html


def test_order_filters(staff_client, staff, customer, products, order):
    other = services.place_order(customer, items(products, nutmeg=1), "collection", "collect", "784", now=NOW)
    services.mark_paid(staff, other, "RCPT-1", now=NOW)
    html = staff_client.get(reverse("staff_orders"), {"status": "paid"}).content.decode()
    assert other.number in html and order.number not in html
    html = staff_client.get(reverse("staff_orders"), {"q": order.number}).content.decode()
    assert order.number in html and other.number not in html


# ---- Stock ---------------------------------------------------------------------------------

def test_receive_delivery_and_count(staff_client, products):
    url = lambda action: reverse("staff_stock_action", args=["BK-03", action])
    html = staff_client.post(url("restock"), {"quantity": 5, "note": "inv 77"}, follow=True).content.decode()
    assert "now has 8 in stock" in html
    html = staff_client.post(url("count"), {"quantity": 6, "note": "two chipped"}, follow=True).content.decode()
    assert "corrected by -2 to 6" in html
    assert "Enter a whole number" in staff_client.post(url("restock"), {"quantity": "x"}, follow=True).content.decode()
    assert Product.objects.get(sku="BK-03").stock == 6


def test_low_stock_list(staff_client, products):
    html = staff_client.get(reverse("staff_stock"), {"low": 1}).content.decode()
    assert "Woven Palm Basket" in html and "Scotch Bonnet" not in html     # 3 left vs 20 left


# ---- The sales and VAT report --------------------------------------------------------------

@pytest.fixture
def month(customer, staff, products, code):
    """Four orders in November with known figures."""
    def sale(when, paid_hours, cancel=False, **qty):
        o = services.place_order(customer, items(products, **qty), "collection", "deliver", "784",
                                 delivery_address="Bay St", now=when)
        if paid_hours is not None:
            services.mark_paid(staff, o, "R", now=when + timedelta(hours=paid_hours))
        if cancel:
            services.cancel_order(staff, o, "damaged in transit", now=when + timedelta(hours=paid_hours + 1))
        return o
    return {
        "a": sale(datetime(2026, 11, 3, 14, 0, tzinfo=tz.utc), 1, hot_sauce=2),            # 23.20 + 25 delivery
        "b": sale(datetime(2026, 11, 3, 15, 0, tzinfo=tz.utc), 2, nutmeg=2, basket=1),     # 17.00 zero-rated + 75
        # placed in October, paid in November: belongs to November
        "c": sale(datetime(2026, 10, 31, 22, 0, tzinfo=tz.utc), 8, hot_sauce=1),
        "cancelled": sale(datetime(2026, 11, 4, 12, 0, tzinfo=tz.utc), 1, cancel=True, basket=1),
        "unpaid": sale(datetime(2026, 11, 5, 12, 0, tzinfo=tz.utc), None, hot_sauce=1),
    }


def test_report_counts_paid_orders_only(month):
    r = reports.sales_report(date(2026, 11, 1), date(2026, 11, 30))
    assert {o.number for o in r.orders} == {month[k].number for k in "abc"}
    assert r.gross_cents == sum(month[k].total_cents for k in "abc")
    assert r.vat_cents == sum(month[k].vat_cents for k in "abc")
    assert [o.number for o in r.refunds_due] == [month["cancelled"].number]


def test_vat_rows_add_up_to_the_totals(month):
    r = reports.sales_report(date(2026, 11, 1), date(2026, 11, 30))
    assert sum(v.gross_cents for v in r.vat_rows) == r.gross_cents
    assert sum(v.vat_cents for v in r.vat_rows) == r.vat_cents
    goods, delivery, zero = r.vat_rows
    assert zero.gross_cents == 1700 and zero.vat_cents == 0
    assert delivery.gross_cents == 3 * 2500 and delivery.vat_cents == 3 * 345


def test_days_have_no_gaps_and_use_local_time(month):
    r = reports.sales_report(date(2026, 11, 1), date(2026, 11, 7))
    assert [d for d, _, _ in r.days] == [date(2026, 11, d) for d in range(1, 8)]
    by_day = {d: g for d, g, _ in r.days}
    # order c was paid at 06:00 UTC on 1 Nov = 02:00 in St Vincent (UTC-4): still 1 Nov
    assert by_day[date(2026, 11, 1)] == month["c"].total_cents
    assert by_day[date(2026, 11, 3)] == month["a"].total_cents + month["b"].total_cents


def test_late_evening_payment_counts_on_the_local_day(customer, staff, products):
    """23:30 in St Vincent is 03:30 UTC the next day; VAT belongs to the local date."""
    o = services.place_order(customer, items(products, nutmeg=1), "collection", "collect", "784",
                             now=datetime(2026, 11, 10, 20, 0, tzinfo=tz.utc))
    services.mark_paid(staff, o, "R", now=datetime(2026, 11, 11, 3, 30, tzinfo=tz.utc))
    assert reports.sales_report(date(2026, 11, 10), date(2026, 11, 10)).order_count == 1
    assert reports.sales_report(date(2026, 11, 11), date(2026, 11, 11)).order_count == 0


def test_top_products(month):
    top = reports.sales_report(date(2026, 11, 1), date(2026, 11, 30)).top_products
    assert top[0][1] == "BK-03"                                    # the basket earns the most
    assert sum(q for _, sku, q, _ in top if sku == "HS-01") == 3


def test_sales_page_and_period_handling(staff_client, month):
    html = staff_client.get(reverse("staff_sales"), {"from": "2026-11-01", "to": "2026-11-30"}).content.decode()
    assert "3 orders" in html and month["cancelled"].number in html     # listed as refund, not as a sale
    r = staff_client.get(reverse("staff_sales"), {"from": "2026-11-30", "to": "2026-11-01"})      # reversed
    assert r.status_code == 200
    assert staff_client.get(reverse("staff_sales"), {"from": "rubbish"}).status_code == 200


# ---- Excel export ----------------------------------------------------------------------

def workbook(month):
    return load_workbook(io.BytesIO(exports.sales_workbook(
        reports.sales_report(date(2026, 11, 1), date(2026, 11, 30)), "Tradewind")))


def test_excel_totals_are_real_formulas(month):
    wb = workbook(month)
    summary, days, orders = wb["Summary"], wb["By day"], wb["Orders"]
    total_row = [r for r in summary.iter_rows(min_row=5) if r[0].value == "Total"][0]
    assert all(c.data_type == "f" and c.value.startswith("=SUM(") for c in total_row[1:4])
    last = days[days.max_row]
    assert last[1].data_type == "f" and last[2].data_type == "f"   # both day totals calculate
    assert orders["I2"].value == "=F2-G2+H2" and orders["I2"].data_type == "f"
    goods = reports.sales_report(date(2026, 11, 1), date(2026, 11, 30)).vat_rows[0]
    assert summary["A5"].value == "Goods at 16%" and summary["B5"].value == goods.gross_cents / 100


def test_excel_never_turns_customer_text_into_formulas(staff, products):
    from django.contrib.auth import get_user_model
    evil = get_user_model().objects.create_user("evil@example.test", "evil@example.test", "x",
                                                first_name='=HYPERLINK("http://evil.example","x")')
    o = services.place_order(evil, items(products, nutmeg=1), "collection", "collect", "784",
                             now=datetime(2026, 11, 3, 14, 0, tzinfo=tz.utc))
    services.mark_paid(staff, o, "=1+1", now=datetime(2026, 11, 3, 15, 0, tzinfo=tz.utc))
    wb = load_workbook(io.BytesIO(exports.sales_workbook(
        reports.sales_report(date(2026, 11, 1), date(2026, 11, 30)), "Tradewind")))
    cells = [c for c in wb["Orders"]["C"][1:] + wb["Orders"]["E"][1:]]
    assert all(c.data_type == "s" for c in cells)


def test_excel_download(staff_client, month):
    r = staff_client.get(reverse("staff_sales_xlsx"), {"from": "2026-11-01", "to": "2026-11-30"})
    assert r.content[:2] == b"PK"
    assert 'filename="tradewind-sales-2026-11-01-to-2026-11-30.xlsx"' in r["Content-Disposition"]


# ---- Dashboard -----------------------------------------------------------------------------

def test_dashboard_counts(staff, customer, products, order):
    paid = services.place_order(customer, items(products, nutmeg=1), "collection", "collect", "784", now=NOW)
    services.mark_paid(staff, paid, "R", now=NOW)
    counts = reports.dashboard_counts(now=NOW + timedelta(days=4))
    assert counts["awaiting"] == 1 and counts["overdue"] == 1 and counts["to_pack"] == 1
    assert counts["low_stock"] == 1                                 # baskets: 2 left, alert at 5
