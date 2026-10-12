"""Sales report as an Excel workbook.

Totals are Excel formulas (SUM over the rows above), so an accountant can
check how each figure is built. Text that came from customers (names,
addresses) is always stored as text, never as a formula, so a name like
=HYPERLINK(...) can't run when the file is opened ("formula injection").
"""

import io
from datetime import timedelta

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

MONEY = '#,##0.00;(#,##0.00);"-"'
HEADER_FILL = PatternFill("solid", fgColor="0E5E6F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
DANGEROUS_START = ("=", "+", "-", "@", "\t", "\r")


def put(ws, row, col, value, font=None, money=False, formula=False):
    """Write one cell. Text is always stored as plain text unless formula=True, which only this
    file's own code passes, so nothing a customer typed can ever become a formula."""
    cell = ws.cell(row=row, column=col)
    cell.value = value / 100 if money and isinstance(value, int) else value     # cents -> dollars
    if money:
        cell.number_format = MONEY
    if isinstance(value, str) and value.startswith(DANGEROUS_START) and not formula:
        cell.data_type = "s"
        cell.quotePrefix = True
    if font:
        cell.font = font
    return cell


def _header(ws, row, names, widths):
    for col, (name, width) in enumerate(zip(names, widths), start=1):
        cell = put(ws, row, col, name, HEADER_FONT)
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def sales_workbook(report, shop_name):
    wb = Workbook()

    # ---- Summary, with the VAT figures a return needs ----
    ws = wb.active
    ws.title = "Summary"
    last = timezone.localtime(report.end).date() - timedelta(days=1)
    put(ws, 1, 1, f"{shop_name}: sales and VAT", Font(bold=True, size=14))
    put(ws, 2, 1, f"Orders paid {timezone.localtime(report.start):%d %b %Y} to {last:%d %b %Y}, "
                  "excluding cancelled orders")
    _header(ws, 4, ["VAT category", "Sales incl. VAT", "VAT", "Sales excl. VAT"], [24, 16, 14, 16])
    for i, v in enumerate(report.vat_rows, start=5):
        put(ws, i, 1, v.label)
        put(ws, i, 2, v.gross_cents, money=True)
        put(ws, i, 3, v.vat_cents, money=True)
        put(ws, i, 4, f"=B{i}-C{i}", money=True, formula=True)
    total_row = 5 + len(report.vat_rows)
    put(ws, total_row, 1, "Total", BOLD)
    for col in "BCD":
        put(ws, total_row, " BCD".index(col) + 1, f"=SUM({col}5:{col}{total_row - 1})", BOLD, money=True,
            formula=True)
    put(ws, total_row + 2, 1, "Orders")
    put(ws, total_row + 2, 2, report.order_count)
    put(ws, total_row + 3, 1, "Discounts given")
    put(ws, total_row + 3, 2, report.discount_cents, money=True)
    put(ws, total_row + 4, 1, "Refunds due")
    put(ws, total_row + 4, 2, sum(o.total_cents for o in report.refunds_due), money=True)

    # ---- Day by day ----
    days = wb.create_sheet("By day")
    _header(days, 1, ["Date", "Orders", "Sales incl. VAT"], [14, 10, 16])
    for i, (day, gross, count) in enumerate(report.days, start=2):
        c = put(days, i, 1, day)
        c.number_format = "ddd dd mmm"
        put(days, i, 2, count)
        put(days, i, 3, gross, money=True)
    end = len(report.days) + 1
    put(days, end + 1, 1, "Total", BOLD)
    put(days, end + 1, 2, f"=SUM(B2:B{end})", BOLD, formula=True)
    put(days, end + 1, 3, f"=SUM(C2:C{end})", BOLD, money=True, formula=True)

    # ---- Products ----
    prods = wb.create_sheet("Top products")
    _header(prods, 1, ["Product", "SKU", "Quantity", "Sales incl. VAT"], [34, 10, 10, 16])
    for i, (name, sku, qty, gross) in enumerate(report.top_products, start=2):
        put(prods, i, 1, name)
        put(prods, i, 2, sku)
        put(prods, i, 3, qty)
        put(prods, i, 4, gross, money=True)

    # ---- Every order ----
    orders = wb.create_sheet("Orders")
    _header(orders, 1, ["Order", "Paid", "Customer", "Payment", "Reference", "Subtotal", "Discount", "Delivery",
                        "Total", "VAT", "Status"], [12, 17, 26, 18, 16, 11, 10, 10, 11, 10, 12])
    for i, o in enumerate(report.orders, start=2):
        put(orders, i, 1, o.number)
        c = put(orders, i, 2, timezone.localtime(o.paid_at).replace(tzinfo=None))    # Excel has no time zones
        c.number_format = "dd mmm yyyy hh:mm"
        put(orders, i, 3, o.customer.get_full_name() or o.customer.email)
        put(orders, i, 4, o.get_payment_method_display())
        put(orders, i, 5, o.payment_reference)
        put(orders, i, 6, o.subtotal_cents, money=True)
        put(orders, i, 7, o.discount_cents, money=True)
        put(orders, i, 8, o.delivery_cents, money=True)
        put(orders, i, 9, f"=F{i}-G{i}+H{i}", money=True, formula=True)   # the same sum the database enforces
        put(orders, i, 10, o.vat_cents, money=True)
        put(orders, i, 11, o.get_status_display())
    orders.auto_filter.ref = f"A1:K{max(len(report.orders) + 1, 2)}"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
